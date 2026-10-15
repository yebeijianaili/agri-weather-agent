from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app import create_app  # noqa: E402


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> int:
    runtime = ROOT / ".runtime"
    exports = ROOT / "exports"
    evidence = ROOT / "docs" / "evidence"
    runtime.mkdir(parents=True, exist_ok=True)
    exports.mkdir(parents=True, exist_ok=True)
    evidence.mkdir(parents=True, exist_ok=True)

    database = runtime / "release_verify.db"
    if database.exists():
        database.unlink()

    api = create_app(database, seed=False, auto_analyze=False)
    sample_path = ROOT / "data" / "sample_weather_2024.csv"
    sample_bytes = sample_path.read_bytes()

    with TestClient(api) as client:
        health_before = client.get("/api/health")
        assert health_before.status_code == 200
        assert health_before.json()["records"] == 0

        imported_response = client.post(
            "/api/observations/import?replace=true",
            files={"file": (sample_path.name, sample_bytes, "text/csv")},
        )
        assert imported_response.status_code == 200, imported_response.text
        imported = imported_response.json()
        # The teaching sample intentionally contains three missing core-solar
        # cells.  A real CSV import must reject those rows while preserving the
        # total-row audit and analysing every accepted row.
        assert imported["total_rows"] == 1098
        assert imported["imported"] >= 1095
        assert imported["imported"] + imported["rejected"] == imported["total_rows"]

        analysis_response = client.post("/api/analysis/run", json={"year": 2024})
        assert analysis_response.status_code == 200, analysis_response.text
        analysis = analysis_response.json()
        analysis_id = analysis["analysis_id"]
        result = analysis["result"]
        assert analysis["status"] == "completed"
        assert len(result["seasonal"]) >= 4
        expected_analysis_rows = imported.get("quality", {}).get("valid_records", imported["imported"])
        assert result["overview"]["record_count"] == expected_analysis_rows

        agent_response = client.post(
            "/api/agent/query",
            json={"analysis_id": analysis_id, "question": "哪个季节降水最多，有干旱风险吗？"},
        )
        assert agent_response.status_code == 200, agent_response.text
        agent = agent_response.json()
        assert agent["analysis_id"] == analysis_id
        assert agent["evidence"]

        csv_response = client.get("/api/export/seasonal.csv", params={"analysis_id": analysis_id})
        md_response = client.get("/api/export/report.md", params={"analysis_id": analysis_id})
        xlsx_response = client.get("/api/export/report.xlsx", params={"analysis_id": analysis_id})
        assert csv_response.status_code == 200
        assert md_response.status_code == 200
        assert xlsx_response.status_code == 200
        csv_bytes = csv_response.content
        md_bytes = md_response.content
        xlsx_bytes = xlsx_response.content
        rows = list(csv.DictReader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
        assert rows and {"analysis_id", "plot_id", "season_year", "season"}.issubset(rows[0])
        assert "# 农田气象时序分析报告" in md_bytes.decode("utf-8")
        workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        assert workbook.sheetnames == ["执行摘要", "季节统计", "风险清单", "方法与口径"]
        assert workbook["执行摘要"]["A1"].value == "农田气象时序分析报告"
        assert workbook["季节统计"]._charts

        csv_path = exports / "seasonal_statistics.csv"
        md_path = exports / "analysis_report.md"
        xlsx_path = exports / "analysis_report.xlsx"
        csv_path.write_bytes(csv_bytes)
        md_path.write_bytes(md_bytes)
        xlsx_path.write_bytes(xlsx_bytes)

        audit_response = client.get("/api/audit/events", params={"limit": 100})
        assert audit_response.status_code == 200
        audit = audit_response.json()

    payload = {
        "verified_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "workflow": ["CSV import", "quality validation", "analysis", "Agent query", "CSV export", "Markdown export", "Excel export", "audit readback"],
        "sample": {
            "path": str(sample_path.relative_to(ROOT)),
            "bytes": len(sample_bytes),
            "sha256": sha256(sample_bytes),
        },
        "import": imported,
        "analysis": {
            "analysis_id": analysis_id,
            "status": analysis["status"],
            "records": result["overview"]["record_count"],
            "plots": result["overview"]["plot_count"],
            "quality_score": result["quality"]["score"],
            "triggered_risks": len([item for item in result["risks"] if item.get("status") == "触发"]),
        },
        "agent": {"answer": agent["answer"], "evidence_count": len(agent["evidence"])},
        "exports": {
            "seasonal_csv": {"path": "exports/seasonal_statistics.csv", "rows": len(rows), "bytes": len(csv_bytes), "sha256": sha256(csv_bytes)},
            "markdown_report": {"path": "exports/analysis_report.md", "bytes": len(md_bytes), "sha256": sha256(md_bytes)},
            "excel_report": {
                "path": "exports/analysis_report.xlsx",
                "sheets": workbook.sheetnames,
                "bytes": len(xlsx_bytes),
                "sha256": sha256(xlsx_bytes),
            },
        },
        "audit_event_count": audit.get("total", len(audit.get("items", []))) if isinstance(audit, dict) else len(audit),
    }
    evidence_path = evidence / "api_verification.json"
    evidence_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print(f"\nVERIFIED_EVIDENCE={evidence_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
