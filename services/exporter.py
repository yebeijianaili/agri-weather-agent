from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4D3A")
SUBHEADER_FILL = PatternFill("solid", fgColor="DDEBDD")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _seasonal_rows(analysis: dict[str, Any]) -> list[dict[str, Any]]:
    """Normalize seasonal results for both CSV and Excel exports."""
    result = analysis["result"]
    source = result.get("seasonal_by_plot") or result.get("seasonal_by_year") or result["seasonal"]
    rows: list[dict[str, Any]] = []
    for item in source:
        rows.append(
            {
                "analysis_id": analysis["analysis_id"],
                "plot_id": item.get("plot_id", "ALL"),
                "season_year": item.get("season_year"),
                "season": item.get("season"),
                "period_start": item.get("period_start", result["overview"].get("date_start")),
                "period_end": item.get("period_end", result["overview"].get("date_end")),
                "expected_days": item.get("expected_days", item.get("day_count")),
                "valid_days": item.get("valid_days", item.get("day_count")),
                "coverage_pct": item.get("coverage_pct", 100),
                "temperature_avg_c": item.get("avg_temperature_c"),
                "temperature_max_c": item.get("max_temperature_c"),
                "temperature_min_c": item.get("min_temperature_c"),
                "temperature_trend_c_30d": item.get("temperature_trend_c_30d"),
                "precipitation_total_mm": item.get("total_precipitation_mm"),
                "rainy_days": item.get("rainy_days"),
                "longest_dry_spell_days": item.get("longest_dry_spell_days"),
                "solar_radiation_avg_wm2": item.get("avg_solar_radiation_wm2"),
                "low_light_days": item.get("low_light_days"),
                "humidity_avg_pct": item.get("avg_humidity_pct"),
                "outlier_count": item.get("outlier_count", 0),
                "risk_count": item.get("risk_count", 0),
            }
        )
    return rows


def seasonal_csv(analysis: dict[str, Any]) -> str:
    output = io.StringIO()
    fields = [
        "analysis_id",
        "plot_id",
        "season_year",
        "season",
        "period_start",
        "period_end",
        "expected_days",
        "valid_days",
        "coverage_pct",
        "temperature_avg_c",
        "temperature_max_c",
        "temperature_min_c",
        "temperature_trend_c_30d",
        "precipitation_total_mm",
        "rainy_days",
        "longest_dry_spell_days",
        "solar_radiation_avg_wm2",
        "low_light_days",
        "humidity_avg_pct",
        "outlier_count",
        "risk_count",
    ]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    rows = _seasonal_rows(analysis)
    writer.writerows(rows)
    return "\ufeff" + output.getvalue()


def excel_report(analysis: dict[str, Any]) -> bytes:
    """Build a styled, auditable Excel report entirely in memory."""
    result = analysis["result"]
    overview = result["overview"]
    workbook = Workbook()

    summary = workbook.active
    summary.title = "执行摘要"
    summary.merge_cells("A1:D1")
    summary["A1"] = "农田气象时序分析报告"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = HEADER_FILL
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.row_dimensions[1].height = 30
    metadata = [
        ("分析任务", f"#{analysis['analysis_id']}"),
        ("生成时间", analysis["created_at"]),
        ("数据范围", f"{overview['date_start']} 至 {overview['date_end']}"),
        ("地块数量", overview["plot_count"]),
        ("记录数量", overview["record_count"]),
        ("平均气温(°C)", overview["avg_temperature_c"]),
        ("累计降水(mm)", overview["total_precipitation_mm"]),
        ("平均太阳辐射(W/m²)", overview["avg_solar_radiation_wm2"]),
        ("平均相对湿度(%)", overview["avg_humidity_pct"]),
        ("入库质量评分", result["quality"]["score"]),
    ]
    for row_index, (label, value) in enumerate(metadata, 3):
        summary.cell(row_index, 1, label).font = Font(bold=True, color="1F4D3A")
        summary.cell(row_index, 2, value)
    narrative_row = len(metadata) + 4
    summary.cell(narrative_row, 1, "执行摘要").font = Font(bold=True, color="1F4D3A")
    summary.merge_cells(start_row=narrative_row + 1, start_column=1, end_row=narrative_row + 4, end_column=4)
    summary.cell(narrative_row + 1, 1, analysis["summary"])
    summary.cell(narrative_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 18

    statistics = workbook.create_sheet("季节统计")
    rows = _seasonal_rows(analysis)
    headers = list(rows[0].keys()) if rows else []
    for column, header in enumerate(headers, 1):
        cell = statistics.cell(1, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            statistics.cell(row_index, column, row.get(header))
    statistics.freeze_panes = "A2"
    if headers:
        statistics.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for column, header in enumerate(headers, 1):
        sample_width = max([len(str(row.get(header, ""))) for row in rows] + [len(header)])
        statistics.column_dimensions[get_column_letter(column)].width = min(max(sample_width + 2, 12), 28)

    if rows and "season" in headers and "temperature_avg_c" in headers:
        chart = BarChart()
        chart.title = "季节平均气温对比"
        chart.y_axis.title = "温度 (°C)"
        chart.x_axis.title = "季节记录"
        data_column = headers.index("temperature_avg_c") + 1
        season_column = headers.index("season") + 1
        chart.add_data(Reference(statistics, min_col=data_column, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(statistics, min_col=season_column, min_row=2, max_row=len(rows) + 1))
        chart.height = 8
        chart.width = 15
        statistics.add_chart(chart, f"{get_column_letter(len(headers) + 2)}2")

    risks = workbook.create_sheet("风险清单")
    risk_headers = ["风险名称", "等级", "状态", "证据", "建议"]
    for column, header in enumerate(risk_headers, 1):
        cell = risks.cell(1, column, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row_index, risk in enumerate(result["risks"], 2):
        values = [risk["name"], risk["severity"], risk["status"], risk["evidence"], risk["recommendation"]]
        for column, value in enumerate(values, 1):
            risks.cell(row_index, column, value).alignment = Alignment(wrap_text=True, vertical="top")
    risks.freeze_panes = "A2"
    for column, width in enumerate([24, 12, 14, 48, 48], 1):
        risks.column_dimensions[get_column_letter(column)].width = width

    methods = workbook.create_sheet("方法与口径")
    methods.append(["项目", "说明"])
    for cell in methods[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    methodology = result["methodology"]
    method_rows = [
        ("季节划分", methodology["season_definition"]),
        ("降水口径", methodology["precipitation_method"]),
        ("趋势方法", methodology["trend_method"]),
        ("异常方法", methodology["anomaly_method"]),
        ("质量评分边界", result["quality"].get("interpretation", "仅评价入库完整性、范围与重复。")),
        ("决策边界", "自动结论用于数据辅助研判，田间管理需结合土壤、作物品种与现场巡查。"),
    ]
    for row in method_rows:
        methods.append(row)
    methods.column_dimensions["A"].width = 22
    methods.column_dimensions["B"].width = 90
    for row in methods.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def markdown_report(analysis: dict[str, Any]) -> str:
    result = analysis["result"]
    overview = result["overview"]
    lines = [
        "# 农田气象时序分析报告",
        "",
        f"- 分析任务：#{analysis['analysis_id']}",
        f"- 生成时间：{analysis['created_at']}",
        f"- 分析状态：{analysis['status']}",
        f"- 数据范围：{overview['date_start']} 至 {overview['date_end']}",
        f"- 地块数量：{overview['plot_count']}",
        f"- 记录数量：{overview['record_count']}",
        "",
        "## 执行摘要",
        "",
        analysis["summary"],
        "",
        "## 核心指标",
        "",
        f"- 平均气温：{overview['avg_temperature_c']} °C（{overview['min_temperature_c']}–{overview['max_temperature_c']} °C）",
        f"- 区域累计降水：{overview['total_precipitation_mm']} mm，雨日 {overview['rainy_days']} 天",
        f"- 平均太阳辐射：{overview['avg_solar_radiation_wm2']} W/m²",
        f"- 平均相对湿度：{overview['avg_humidity_pct']}%",
        f"- 最长连续少雨：{overview['longest_dry_spell_days']} 天",
        "",
        "## 四季对比",
        "",
        "| 季节 | 平均温度(°C) | 累计降水(mm) | 雨日 | 平均太阳辐射(W/m²) | 低光照日 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for season in result["seasonal"]:
        lines.append(f"| {season['season']} | {season['avg_temperature_c']} | {season['total_precipitation_mm']} | {season['rainy_days']} | {season['avg_solar_radiation_wm2']} | {season['low_light_days']} |")
    lines.extend(["", "## Agent 结论", ""])
    lines.extend(f"{index}. {value}" for index, value in enumerate(result["conclusions"], 1))
    lines.extend(["", "## 风险清单", ""])
    for risk in result["risks"]:
        lines.append(f"- **{risk['name']} / {risk['severity']} / {risk['status']}**：{risk['evidence']}；{risk['recommendation']}")
    quality = result["quality"]
    lines.extend(
        [
            "",
            "## 入库质量与方法",
            "",
            f"- 入库质量评分：{quality['score']}（{quality['grade']}）",
            f"- 评分口径：{quality.get('interpretation', '仅评价入库完整性、范围与重复，不代表测量准确性。')}",
            f"- 缺测：{quality['missing_values']}；疑似记录：{quality['suspect_records']}；重复：{quality['duplicate_records']}",
            f"- 季节划分：{result['methodology']['season_definition']}",
            f"- 降水口径：{result['methodology']['precipitation_method']}",
            f"- 趋势方法：{result['methodology']['trend_method']}",
            f"- 异常方法：{result['methodology']['anomaly_method']}",
            "",
            "> 本报告由规则分析 Agent 基于当前数据自动生成，田间管理决策需结合土壤、作物品种与现场巡查。",
            "",
        ]
    )
    return "\n".join(lines)
