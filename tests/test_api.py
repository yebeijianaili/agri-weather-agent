from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

import pytest


def test_health_and_seeded_dataset(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert payload["version"] == "1.2.0"
    assert payload["records"] == 1098
    assert payload["agent_runtime"]["mode"] == "rule_fallback"
    assert payload["agent_runtime"]["provider"] == "rules"
    assert payload["agent_runtime"]["fallback_used"] is True

    source = client.get("/api/data/source")
    assert source.status_code == 200
    assert source.json()["source"] is None

    overview = client.get("/api/overview").json()
    assert overview["dataset"]["plot_count"] == 3
    assert overview["dataset"]["date_start"] == "2024-01-01"
    assert overview["dataset"]["date_end"] == "2024-12-31"
    assert overview["quality"]["missing_values"] == 5
    assert overview["quality"]["score"] > 95
    assert overview["quality"]["scope"] == "ingestion_quality"
    assert overview["quality"]["score_label"] == "入库质量得分"
    assert "不代表" in overview["quality"]["interpretation"]


def test_observation_filters_and_pagination(client):
    plots = client.get("/api/overview").json()["dataset"]["plots"]
    response = client.get("/api/observations", params={"plot_id": plots[0], "year": 2024, "limit": 10, "offset": 3})
    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 366
    assert len(payload["items"]) == 10
    assert all(item["plot_id"] == plots[0] for item in payload["items"])


def test_analysis_is_complete_and_evidence_backed(client):
    response = client.post("/api/analysis/run", json={"year": 2024})
    assert response.status_code == 200
    payload = response.json()
    result = payload["result"]
    assert payload["status"] == "completed"
    assert len(result["seasonal"]) == 4
    assert {item["season"] for item in result["seasonal"]} == {"春季", "夏季", "秋季", "冬季"}
    assert len(result["monthly"]) == 12
    assert result["data_scope"]["source_records"] == 1098
    assert result["data_scope"]["analyzed_records"] == 1093
    assert result["data_scope"]["excluded_non_valid_records"] == 5
    assert set(result["trends"]) == {"temperature_c", "precipitation_mm", "solar_radiation_wm2", "humidity_pct"}
    assert len(result["risks"]) == 7
    assert all("evidence" in risk and "threshold" in risk for risk in result["risks"])
    assert len(result["seasonal_by_year"]) == 5
    december_winter = next(item for item in result["seasonal_by_year"] if item["season"] == "冬季" and item["season_year"] == 2025)
    assert december_winter["day_count"] == 31
    assert {event["code"] for event in result["risk_events"]} >= {
        "HEAT_DAY", "HEAT_STREAK", "COLD_DAY", "HEAVY_RAIN", "DRY_SPELL", "LOW_LIGHT_STREAK", "TEMP_SWING"
    }
    assert all(
        {"plot_id", "start_date", "end_date", "duration_days", "evidence", "threshold"} <= set(event)
        for event in result["risk_events"]
    )
    assert len(result["conclusions"]) >= 4

    detail = client.get(f"/api/analysis/{payload['analysis_id']}")
    assert detail.status_code == 200
    assert detail.json()["analysis_id"] == payload["analysis_id"]
    history = client.get("/api/analysis/history").json()
    assert history["total"] >= 2


def test_agent_query_returns_citations_to_metrics(client):
    response = client.post("/api/agent/query", json={"question": "哪个季节降水最多，有干旱风险吗？", "year": 2024})
    assert response.status_code == 200
    payload = response.json()
    assert "累计" in payload["answer"]
    assert len(payload["evidence"]) >= 3
    assert payload["analysis_id"] >= 1
    assert payload["query_id"] >= 1
    assert payload["context"]["date_start"] == "2024-01-01"
    assert payload["agent_runtime"] == payload["context"]["agent_runtime"]
    assert payload["agent_runtime"]["fallback_reason"] == "provider_disabled"
    assert payload["disclaimer"].startswith("Agent 的结论")
    assert "规则 Agent" not in payload["disclaimer"]
    assert "分析 #" in payload["answer"]
    history = client.get("/api/agent/history").json()
    assert history["total"] == 1
    assert history["items"][0]["context"]["agent_runtime"] == payload["agent_runtime"]


def test_agent_greeting_is_brief_human_and_not_wrapped_as_analysis(client):
    response = client.post("/api/agent/query", json={"question": "你好"})

    assert response.status_code == 200
    payload = response.json()
    assert payload["intent"] == "greeting"
    assert payload["evidence"] == []
    assert "基于分析 #" not in payload["answer"]
    assert "你好" in payload["answer"]
    assert len(payload["answer"]) <= 90
    assert payload["context"]["analysis_id"] == payload["analysis_id"]
    assert payload["suggestions"]


def test_agent_handles_common_conversation_intents_without_fake_evidence(client):
    cases = {
        "你能做什么": "capability",
        "谢谢你": "thanks",
        "再见": "farewell",
    }
    for question, expected_intent in cases.items():
        payload = client.post("/api/agent/query", json={"question": question}).json()
        assert payload["intent"] == expected_intent
        assert payload["evidence"] == []
        assert "基于分析 #" not in payload["answer"]
        assert len(payload["answer"]) <= 90

    domain_payload = client.post(
        "/api/agent/query",
        json={"question": "你好，哪个季节降水最多？"},
    ).json()
    assert domain_payload["intent"] == "precipitation"
    assert domain_payload["evidence"]
    assert "基于分析 #" in domain_payload["answer"]


def test_official_dataset_source_is_persisted_and_exposed(client):
    provenance = {
        "source_key": "nasa-power-test",
        "provider": "NASA Langley Research Center POWER Project",
        "dataset_name": "NASA/POWER Source Native Resolution Daily Data",
        "landing_page": "https://power.larc.nasa.gov/",
        "query_urls": ["https://power.larc.nasa.gov/api/temporal/daily/point?test=1"],
        "retrieved_at": "2026-08-22T18:00:00+08:00",
        "api_version": "v2.9.7",
    }
    saved = client.app.state.repository.save_dataset_source(provenance)

    response = client.get("/api/data/source")
    assert response.status_code == 200
    source = response.json()["source"]
    assert source["source_id"] == saved["source_id"]
    assert source["provider"].startswith("NASA")
    assert source["provenance"]["api_version"] == "v2.9.7"
    assert client.get("/api/overview").json()["dataset"]["source"]["source_key"] == "nasa-power-test"


def test_dataset_changing_user_import_does_not_keep_nasa_as_active_source(client):
    repository = client.app.state.repository
    repository.save_dataset_source(
        {
            "source_key": "nasa-before-user-import",
            "provider": "NASA Langley Research Center POWER Project",
            "dataset_name": "NASA/POWER Source Native Resolution Daily Data",
            "landing_page": "https://power.larc.nasa.gov/",
            "query_urls": ["https://power.larc.nasa.gov/api/temporal/daily/point?test=1"],
            "retrieved_at": "2026-08-22T18:00:00+08:00",
        }
    )
    appended_csv = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2,humidity_pct\n"
        "2026-01-01,USER-PLOT,20,1.2,220,70\n"
    ).encode("utf-8")

    appended = client.post(
        "/api/observations/import",
        files={"file": ("user-append.csv", appended_csv, "text/csv")},
    )

    assert appended.status_code == 200
    appended_source = appended.json()["data_source"]
    assert appended_source["provider"] == "混合来源数据集"
    assert appended_source["provenance"]["source_type"] == "mixed_dataset"
    assert appended_source["provenance"]["parent_source"]["provider"].startswith("NASA")
    assert client.get("/api/data/source").json()["source"]["source_key"] == appended_source["source_key"]

    replacement_csv = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2,humidity_pct\n"
        "2027-01-01,USER-ONLY,21,0,230,68\n"
    ).encode("utf-8")
    replaced = client.post(
        "/api/observations/import?replace=true",
        files={"file": ("user-replace.csv", replacement_csv, "text/csv")},
    )

    assert replaced.status_code == 200
    replacement_source = replaced.json()["data_source"]
    assert replacement_source["provider"] == "用户上传 CSV"
    assert replacement_source["provenance"]["source_type"] == "user_csv"
    assert replacement_source["provenance"]["parent_source"] is None
    assert client.get("/api/data/source").json()["source"]["source_key"] == replacement_source["source_key"]


def test_changed_dataset_drops_stale_source_if_new_source_registration_fails(client, monkeypatch):
    repository = client.app.state.repository
    repository.save_dataset_source(
        {
            "source_key": "nasa-before-registration-failure",
            "provider": "NASA Langley Research Center POWER Project",
            "dataset_name": "NASA/POWER Source Native Resolution Daily Data",
            "landing_page": "https://power.larc.nasa.gov/",
            "query_urls": [],
            "retrieved_at": "2026-08-22T18:00:00+08:00",
        }
    )
    monkeypatch.setattr(repository, "save_dataset_source", lambda _payload: (_ for _ in ()).throw(RuntimeError("source write failed")))
    csv_data = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2,humidity_pct\n"
        "2026-01-01,USER-PLOT,20,1.2,220,70\n"
    ).encode("utf-8")

    with pytest.raises(RuntimeError, match="source write failed"):
        client.post(
            "/api/observations/import",
            files={"file": ("user.csv", csv_data, "text/csv")},
        )

    assert repository.observation_count() == 1099
    assert repository.active_dataset_source() is None


def test_csv_import_updates_inserts_and_rejects(client):
    csv_data = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2,humidity_pct,quality_flag\n"
        "2024-01-01,A-01 北区麦田,21.5,0,300,66,valid\n"
        "2025-01-01,D-04 示范田,18.2,2.1,310,71,valid\n"
        "bad-date,D-04 示范田,18.2,2.1,310,71,valid\n"
    ).encode("utf-8")
    response = client.post("/api/observations/import", files={"file": ("weather.csv", csv_data, "text/csv")})
    assert response.status_code == 200
    payload = response.json()
    assert payload["imported"] == 1
    assert payload["updated"] == 0
    assert payload["rejected"] == 2
    assert payload["conflict_duplicates"] == 1
    assert payload["reason_counts"]["duplicate_conflict"] == 1
    assert payload["import_id"] >= 1
    assert payload["filename"] == "weather.csv"
    assert len(payload["sha256"]) == 64
    assert payload["quality"]["total_records"] == 1099

    history = client.get("/api/imports/history").json()
    assert history["total"] == 1
    assert history["items"][0]["import_id"] == payload["import_id"]


def test_exports_are_real_downloads(client):
    csv_response = client.get("/api/export/seasonal.csv")
    assert csv_response.status_code == 200
    assert "text/csv" in csv_response.headers["content-type"]
    reader = csv.DictReader(io.StringIO(csv_response.text.lstrip("\ufeff")))
    expected_fields = {
        "analysis_id", "plot_id", "season_year", "season", "period_start", "period_end",
        "expected_days", "valid_days", "coverage_pct", "temperature_avg_c",
        "temperature_max_c", "temperature_min_c", "temperature_trend_c_30d",
        "precipitation_total_mm", "rainy_days", "longest_dry_spell_days",
        "solar_radiation_avg_wm2", "low_light_days", "humidity_avg_pct",
        "outlier_count", "risk_count",
    }
    assert set(reader.fieldnames or []) == expected_fields
    exported_rows = list(reader)
    assert len(exported_rows) == 15
    assert {row["plot_id"] for row in exported_rows} == {"A-01 北区麦田", "B-02 南区水稻田", "C-03 东区试验田"}
    assert any(row["season"] == "春季" for row in exported_rows)
    assert int(csv_response.headers["x-audit-id"]) >= 1

    md_response = client.get("/api/export/report.md")
    assert md_response.status_code == 200
    assert "# 农田气象时序分析报告" in md_response.text
    assert "## 四季对比" in md_response.text

    xlsx_response = client.get("/api/export/report.xlsx")
    assert xlsx_response.status_code == 200
    assert xlsx_response.content[:2] == b"PK"
    assert "spreadsheetml.sheet" in xlsx_response.headers["content-type"]
    workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=True)
    assert workbook.sheetnames == ["执行摘要", "季节统计", "风险清单", "方法与口径"]
    assert workbook["执行摘要"]["A1"].value == "农田气象时序分析报告"
    assert workbook["季节统计"].max_row == 16
    assert workbook["季节统计"]._charts
    assert int(xlsx_response.headers["x-audit-id"]) >= 1
    audit = client.get("/api/audit/events").json()
    assert audit["total"] == 3


def test_empty_filter_and_invalid_date_are_handled(client):
    missing = client.post("/api/analysis/run", json={"year": 2030})
    assert missing.status_code == 404
    assert missing.json()["detail"]["status"] == "failed"
    failed_id = missing.json()["detail"]["analysis_id"]
    failed_detail = client.get(f"/api/analysis/{failed_id}").json()
    assert failed_detail["status"] == "failed"
    assert client.get("/api/analysis/latest").json()["status"] == "completed"
    invalid = client.get("/api/overview", params={"start_date": "2024-12-01", "end_date": "2024-01-01"})
    assert invalid.status_code == 422


def test_import_requires_core_headers_and_audits_failure(client):
    bad = b"timestamp,plot_id,temperature_c\n2025-01-01,P-1,20\n"
    response = client.post("/api/observations/import", files={"file": ("missing.csv", bad, "text/csv")})
    assert response.status_code == 422
    assert response.json()["detail"]["import_id"] >= 1
    history = client.get("/api/imports/history").json()
    assert history["items"][0]["reason_counts"] == {"file_structure": 1}


def test_import_rejects_blank_required_weather_values(client):
    data = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2,humidity_pct\n"
        "2025-02-01,P-NEW,,1.2,220,70\n"
        "2025-02-02,P-NEW,18,,220,70\n"
        "2025-02-03,P-NEW,18,1.2,,70\n"
    ).encode("utf-8")
    response = client.post("/api/observations/import", files={"file": ("blank-required.csv", data, "text/csv")})
    assert response.status_code == 200
    payload = response.json()
    assert payload["imported"] == 0
    assert payload["rejected"] == 3
    assert payload["reason_counts"] == {"missing_required": 3}


def test_optional_humidity_can_be_missing_and_rows_remain_analyzable(client):
    data = (
        "timestamp,plot_id,temperature_c,precipitation_mm,solar_radiation_wm2\n"
        "2025-03-01,P-FIVE,18.5,0,210\n"
        "2025-03-02,P-FIVE,19.5,1.2,220\n"
    ).encode("utf-8")
    imported = client.post(
        "/api/observations/import?replace=true",
        files={"file": ("five-columns.csv", data, "text/csv")},
    )
    assert imported.status_code == 200
    payload = imported.json()
    assert payload["imported"] == 2
    assert payload["rejected"] == 0
    assert payload["warning_counts"] == {"missing_value": 2}
    assert payload["quality"]["missing_by_field"]["humidity_pct"] == 2
    assert payload["quality"]["valid_records"] == 2

    analyzed = client.post("/api/analysis/run", json={"plot_id": "P-FIVE", "year": 2025})
    assert analyzed.status_code == 200
    result = analyzed.json()["result"]
    assert result["data_scope"]["analyzed_records"] == 2
    assert result["overview"]["avg_temperature_c"] == 19.0
    assert result["overview"]["avg_humidity_pct"] is None


def test_failed_analysis_cannot_be_queried_or_exported(client):
    failed = client.post("/api/analysis/run", json={"year": 2031}).json()["detail"]
    analysis_id = failed["analysis_id"]
    query = client.post("/api/agent/query", json={"analysis_id": analysis_id, "question": "有什么风险？"})
    assert query.status_code == 409
    exported = client.get("/api/export/report.md", params={"analysis_id": analysis_id})
    assert exported.status_code == 409


def test_explicit_missing_analysis_id_returns_404(client):
    response = client.post(
        "/api/agent/query",
        json={"analysis_id": 999999, "question": "这批数据有什么风险？", "year": 2024},
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "分析任务不存在"


def test_season_and_cross_year_season_filters(client):
    summer = client.get("/api/overview", params={"year": 2024, "season": "summer"})
    assert summer.status_code == 200
    payload = summer.json()
    assert payload["filters"]["season"] == "夏季"
    assert payload["dataset"]["total_records"] == 92 * 3
    assert payload["dataset"]["date_start"] == "2024-06-01"
    assert payload["dataset"]["date_end"] == "2024-08-31"

    winter_2025 = client.get("/api/observations", params={"season": "冬季", "season_year": 2025, "limit": 500})
    assert winter_2025.status_code == 200
    assert winter_2025.json()["total"] == 31 * 3
    assert all(item["timestamp"].startswith("2024-12") for item in winter_2025.json()["items"])

    analysis = client.post("/api/analysis/run", json={"year": 2024, "season": "autumn"})
    assert analysis.status_code == 200
    assert analysis.json()["filters"]["season"] == "秋季"
    invalid = client.get("/api/overview", params={"season": "monsoon"})
    assert invalid.status_code == 422


def test_filtered_export_does_not_create_analysis_history(client):
    before = client.get("/api/analysis/history").json()["total"]
    exported = client.get("/api/export/seasonal.csv", params={"year": 2024, "season": "spring"})
    assert exported.status_code == 200
    assert "seasonal-analysis-filtered.csv" in exported.headers["content-disposition"]
    after = client.get("/api/analysis/history").json()["total"]
    assert after == before
